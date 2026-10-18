"""Build the Bloc 3 dossier technique (PDF + DOCX) from certification/bloc3-2-dossier-technique.md.

Look and feel follow the Google Doc charte (Calibri body / Play headings, navy #1e2761, coral #d9634a).
Outputs go to certification/export/ (gitignored).

Setup (from the repo root):
    python3 -m venv .venv && .venv/bin/pip install pypandoc_binary python-docx pillow
Also needs: weasyprint, LibreOffice (soffice + its python3-uno), and mermaid-cli (npx @mermaid-js/mermaid-cli).
Run:
    .venv/bin/python certification/tools/build_dossier.py
"""
import base64
import glob
import hashlib
import html
import pathlib
import re
import shutil
import subprocess

import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

TOOLS = pathlib.Path(__file__).resolve().parent
CERT = TOOLS.parent
OUT = CERT / "export"
BUILD = OUT / ".build"
SOURCE = CERT / "bloc3-2-dossier-technique.md"
GLOSSARY = CERT / "glossaire.md"
BASENAME = "dossier-technique-arnaud-thery"
REPO = "https://github.com/St4r4x/restaurant-analytics/blob/main/"
COVER_IMAGE = CERT / "screenshots" / "01-landing.png"
MAX_W_CM, MAX_H_CM = 16.0, 18.0
SCREENSHOT_MAX_H_CM = 11.0  # screenshots illustrate the guide, two fit on a page
PANDOC_FORMAT = "markdown+lists_without_preceding_blankline-tex_math_dollars-subscript-superscript"

# Visual identity taken from the Google Doc "dossier-synthese-soutenance" (exported docx styles)
NAVY, CORAL, GREY, LINE, ZEBRA = "1e2761", "d9634a", "5b6472", "dce3f0", "eef2fb"
BODY_FONT, HEAD_FONT = "Calibri", "Play"

METADATA = '---\nlang: fr-FR\ntoc-title: "Sommaire"\n---\n'
COVER = [
    ("label", "CERTIFICATION RNCP NIVEAU 7"),
    ("sublabel", "Expert en informatique et système d'information"),
    ("title", "Dossier technique — Application développée"),
    ("project", "restaurant-analytics"),
    ("meta", "Candidat : Arnaud Thery"),
    ("meta", "Bloc évalué : Bloc 3 — Concevoir et développer une application informatique"),
    ("meta", "Certificateur : 3W Academy · Formation dispensée via Aflokkat"),
    ("meta", "Code source : github.com/St4r4x/restaurant-analytics"),
    ("meta-last", "Démo en production : restaurant-app-production-3b11.up.railway.app"),
    ("intro", "Ce document est le dossier technique de l'application restaurant-analytics, livrable de "
              "l'épreuve finale portant sur la compétence « Développement d'une application informatique ». "
              "Il présente les fonctionnalités livrées, la revue de code, la stratégie de tests, les procédures "
              "d'installation et de déploiement continu, le plan de sécurité et le plan de maintenance, "
              "conformément au référentiel de certification. Version de référence : 2.4.3, état au 02/10/2026."),
]


# ── Markdown preparation ─────────────────────────────────────────────────────

def mmdc_command() -> list:
    cached = sorted(glob.glob(str(pathlib.Path.home() / ".npm/_npx/*/node_modules/.bin/mmdc")))
    if shutil.which("mmdc"):
        return ["mmdc"]
    if cached:
        return [cached[0]]
    return ["npx", "--yes", "-p", "@mermaid-js/mermaid-cli", "mmdc"]


def width_attr(png: pathlib.Path, scale: int = 1, max_h: float = MAX_H_CM) -> str:
    # scale = mermaid render scale; never enlarge a diagram beyond its natural size (96 dpi)
    w, h = Image.open(png).size
    natural_cm = w / scale * 2.54 / 96
    return f"{{width={min(MAX_W_CM, max_h * w / h, natural_cm):.1f}cm}}"


def render_mermaid(src: str) -> pathlib.Path:
    png = BUILD / f"mmd-{hashlib.sha1(src.encode()).hexdigest()[:10]}.png"
    if not png.exists():
        mmd = png.with_suffix(".mmd")
        mmd.write_text(src)
        subprocess.run(mmdc_command() + ["-i", str(mmd), "-o", str(png), "-s", "3", "-b", "white", "-q"],
                       check=True, capture_output=True, timeout=300)
    return png


def breakable_code(line: str) -> str:
    # long identifiers in table cells (e.g. AuthServiceTest.register_throws_whenPasswordTooWeak) may wrap
    # after "." and "_" instead of letter by letter, so their column keeps a readable width
    if not line.startswith("|"):
        return line
    return re.sub(r"`[^`]+`", lambda m: re.sub(r"([._])", "\\1\u200b", m.group(0)), line)


def fix_links(line: str) -> str:
    # repo-relative links (../docs/x.md) become GitHub links in the exported document
    return re.sub(r"(?<!!)\[([^\]]+)\]\(\.\./([^)\s]+)\)", lambda m: f"[{m.group(1)}]({REPO}{m.group(2)})", line)


def convert_source() -> str:
    """Drop the dossier title and metadata (the cover carries them), promote headings one level."""
    out, in_code, lang, buf, started = [], False, "", [], False
    for line in SOURCE.read_text().splitlines():
        fence = re.match(r"^```(\w*)", line)
        if fence and not in_code:
            in_code, lang, buf = True, fence.group(1), []
            if lang != "mermaid":
                out.append(line)
            continue
        if in_code:
            if line.startswith("```"):
                in_code = False
                if lang == "mermaid":
                    png = render_mermaid("\n".join(buf) + "\n")
                    out += ["", f"![]({png}){width_attr(png, scale=3)}", ""]
                else:
                    out.append(line)
            else:
                (buf if lang == "mermaid" else out).append(line)
            continue
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            started = True
        if not started or line.strip() == "---":
            continue
        if line.startswith("##"):
            line = line[1:]
        img = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", line)
        if img:
            png = CERT / img.group(2)
            line = f"![{img.group(1)}]({png}){width_attr(png, max_h=SCREENSHOT_MAX_H_CM)}"
        out.append(breakable_code(fix_links(line)))
    return "\n".join(out) + "\n"


def filtered_glossary(body: str) -> str:
    """Keep only the glossary rows whose term appears in the dossier."""
    keep = ["# Annexe B — Glossaire", ""]
    for line in GLOSSARY.read_text().splitlines():
        row = re.match(r"^\| \*\*(.+?)\*\* \|", line)
        if not row:
            if line.startswith("|") or not line.startswith("#"):
                keep.append(line)
            continue
        terms = re.split(r" / |, | \(|\)", row.group(1))
        if any(t.strip() and t.strip().lower() in body.lower() for t in terms):
            keep.append(line)
    return "\n".join(keep) + "\n"


def column_widths_cm(rows: list, width_cm: float) -> list:
    """Column widths following each column's longest cell, never narrower than its longest unbreakable
    word (e.g. "CT01").

    Each cell is a (text, code) pair: prose hyphenates, while inline code only wraps at the zero-width
    spaces added by breakable_code() and after "/", and header labels never wrap. Shared by the PDF and
    the DOCX."""
    cols = len(rows[0])
    biggest, longest = [0] * cols, [0] * cols
    for ri, row in enumerate(rows):
        for i, (text, code) in enumerate(row[:cols]):
            biggest[i] = max(biggest[i], min(len(text), 60))
            words = [min(len(w), 13) for w in text.split()]
            words += [min(len(w), 24) for w in re.split(r"[\s\u200b/]+", text if ri == 0 else code)]
            longest[i] = max([longest[i]] + words)
    weights = [max(n, 4) ** 0.75 for n in biggest]
    minimum = [0.16 * n + 0.5 for n in longest]
    big = [i for i in range(cols) if minimum[i] > 2]
    excess = sum(minimum) - width_cm
    if excess > 0:  # squeeze the wide columns, never the narrow ones ("CT01", "Réussi")
        share = sum(minimum[i] for i in big)
        for i in big:
            minimum[i] -= excess * minimum[i] / share
    cm = [width_cm * w / sum(weights) for w in weights]
    for _ in range(cols):
        if not any(cm[i] < minimum[i] - 1e-6 for i in range(cols)):
            break
        for i in range(cols):
            cm[i] = max(cm[i], minimum[i])
        free = [i for i in range(cols) if cm[i] > minimum[i] + 1e-6]
        spare = width_cm - sum(cm[i] for i in range(cols) if i not in free)
        for i in free:
            cm[i] = spare * weights[i] / sum(weights[j] for j in free)
    return cm


def fix_table_layout(page: pathlib.Path):
    # WeasyPrint's auto layout can give a column of code identifiers its full width and push the table
    # off the page: fix every table's column widths instead
    def cell(raw: str) -> tuple:
        plain = lambda part: html.unescape(re.sub(r"<[^>]+>", "", part))
        return plain(raw), " ".join(plain(c) for c in re.findall(r"<code[^>]*>(.*?)</code>", raw, re.S))

    def fix(match: re.Match) -> str:
        table = re.sub(r"<colgroup>.*?</colgroup>", "", match.group(0), flags=re.S)
        rows = [[cell(c) for c in re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", tr, re.S)]
                for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", table, re.S)]
        cm = column_widths_cm(rows, 16.5)
        cols = "".join(f'<col style="width:{100 * c / sum(cm):.1f}%">' for c in cm)
        return re.sub(r"<table([^>]*)>", rf'<table\1 style="table-layout:fixed"><colgroup>{cols}</colgroup>',
                      table, count=1)

    page.write_text(re.sub(r"<table.*?</table>", fix, page.read_text(), flags=re.S))


# ── DOCX styling ─────────────────────────────────────────────────────────────

def _set_font(style_or_run, name: str):
    style_or_run.font.name = name
    rpr = style_or_run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    # theme fonts (pandoc's reference.docx uses Aptos) take precedence over explicit ones
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(attr), None)


def _el(tag: str, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def _replace(parent, tag: str, new):
    for old in parent.findall(qn(tag)):
        parent.remove(old)
    parent.append(new)


def style_reference_docx(path: pathlib.Path):
    doc = Document(str(path))
    st = doc.styles
    names = {s.name for s in st}
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in names:
            _set_font(st[name], BODY_FONT)
            st[name].font.size = Pt(10.5)
    for name in ("Body Text", "First Paragraph"):
        fmt = st[name].paragraph_format
        fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(0), Pt(8), 1.0
    st["Compact"].paragraph_format.space_after = Pt(4.5)
    headings = {  # level: (font, size pt, color, space before, space after)
        1: (HEAD_FONT, 16, NAVY, 18, 8), 2: (HEAD_FONT, 13, NAVY, 13, 6),
        3: (BODY_FONT, 11.5, CORAL, 10, 5), 4: (BODY_FONT, 10.5, NAVY, 8, 4),
    }
    for lvl, (font, size, color, before, after) in headings.items():
        s = st[f"Heading {lvl}"]
        _set_font(s, font)
        s.font.size, s.font.bold, s.font.italic = Pt(size), True, False
        s.font.color.rgb = RGBColor.from_string(color.upper())
        s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(before), Pt(after)
    st["Heading 1"].paragraph_format.page_break_before = True
    if "TOC Heading" in names:
        _set_font(st["TOC Heading"], HEAD_FONT)
        st["TOC Heading"].font.size, st["TOC Heading"].font.bold = Pt(16), True
        st["TOC Heading"].font.color.rgb = RGBColor.from_string(NAVY.upper())
    if "Hyperlink" in names:
        st["Hyperlink"].font.color.rgb = RGBColor.from_string(NAVY.upper())
    for name in ("Source Code", "Verbatim Char"):
        if name in names:
            st[name].font.size = Pt(8.5)
    for sec in doc.sections:  # A4, margins of the Google Doc (0.75 in / 0.875 in)
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(1.905)
        sec.left_margin = sec.right_margin = Cm(2.2225)
    doc.save(str(path))


def style_docx_tables(path: pathlib.Path, text_width_cm: float = 16.5):
    doc = Document(str(path))
    for table in doc.tables:
        cols = len(table.columns)
        rows = [[(cell.text, " ".join(r.text for p in cell.paragraphs for r in p.runs
                                      if r.style is not None and r.style.name == "Verbatim Char"))
                 for cell in row.cells[:cols]] for row in table.rows]
        widths = [Cm(c) for c in column_widths_cm(rows, text_width_cm)]
        tblPr = table._tbl.tblPr
        _replace(tblPr, "w:tblW", _el("w:tblW", w=int(Cm(text_width_cm).twips), type="dxa"))
        borders = _el("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borders.append(_el(f"w:{edge}", val="single", sz=2, space=0, color=LINE))
        _replace(tblPr, "w:tblBorders", borders)
        _replace(tblPr, "w:tblLayout", _el("w:tblLayout", type="fixed"))
        _replace(tblPr, "w:jc", _el("w:jc", val="center"))
        margins = _el("w:tblCellMar")
        for side, w in (("top", 80), ("left", 100), ("bottom", 80), ("right", 100)):
            margins.append(_el(f"w:{side}", w=w, type="dxa"))
        _replace(tblPr, "w:tblCellMar", margins)
        for i, grid_col in enumerate(table._tbl.tblGrid.findall(qn("w:gridCol"))):
            grid_col.set(qn("w:w"), str(int(widths[i].twips)))
        for ri, row in enumerate(table.rows):
            header = ri == 0
            # a row never splits across pages (same as tr { break-inside: avoid } in the PDF)
            _replace(row._tr.get_or_add_trPr(), "w:cantSplit", _el("w:cantSplit"))
            fill = NAVY if header else ("ffffff" if ri % 2 else ZEBRA)
            for i, cell in enumerate(row.cells[:cols]):
                cell.width = widths[i]
                tcPr = cell._tc.get_or_add_tcPr()
                _replace(tcPr, "w:shd", _el("w:shd", val="clear", color="auto", fill=fill))
                _replace(tcPr, "w:vAlign", _el("w:vAlign", val="center"))
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        r.font.size = Pt(10 if header else 9.5)
                        if header:
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    doc.save(str(path))


def center_docx_images(path: pathlib.Path):
    doc = Document(str(path))
    for p in doc.paragraphs:
        if p._p.findall(".//" + qn("w:drawing")):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(str(path))


def add_docx_cover(path: pathlib.Path):
    # Same cover layout as the Google Doc: grey label, navy title, coral project name, meta lines, intro, screenshot
    looks = {  # kind: (size, bold, italic, color, space before, space after)
        "label": (11, True, False, GREY, 20, 5), "sublabel": (11, False, False, GREY, 0, 5),
        "title": (22, True, False, NAVY, 15, 5), "project": (16, False, True, CORAL, 0, 15),
        "meta": (11, False, False, None, 0, 3), "meta-last": (11, False, False, None, 0, 20),
        "intro": (10.5, False, True, GREY, 0, 8),
    }
    doc = Document(str(path))
    first = doc.element.body[0]
    for kind, text in COVER:
        size, bold, italic, color, before, after = looks[kind]
        p = doc.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        r = p.add_run(text)
        _set_font(r, BODY_FONT)
        r.font.size, r.font.bold, r.font.italic = Pt(size), bold, italic
        if color:
            r.font.color.rgb = RGBColor.from_string(color.upper())
        first.addprevious(p._p)
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(COVER_IMAGE), width=Cm(13.5))
    first.addprevious(pic._p)
    page_break = doc.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)
    first.addprevious(page_break._p)
    doc.save(str(path))


def cover_html() -> str:
    items = "\n".join(f'<p class="cover-{kind}">{html.escape(text)}</p>' for kind, text in COVER)
    data = base64.b64encode(COVER_IMAGE.read_bytes()).decode()
    return f'<section class="cover">\n{items}\n<img src="data:image/png;base64,{data}" alt="">\n</section>\n'


# ── Build ────────────────────────────────────────────────────────────────────

def main():
    BUILD.mkdir(parents=True, exist_ok=True)
    body = convert_source()
    md = BUILD / f"{BASENAME}.md"
    md.write_text(METADATA + body + "\n" + filtered_glossary(body))

    # PDF: standalone HTML + print stylesheet, rendered by WeasyPrint
    cover = BUILD / "cover.html"
    cover.write_text(cover_html())
    page = BUILD / f"{BASENAME}.html"
    pypandoc.convert_file(str(md), "html5", format=PANDOC_FORMAT, outputfile=str(page),
                          extra_args=["--standalone", "--toc", "--toc-depth=2", "--columns=100000",
                                      f"--css={TOOLS / 'print.css'}", f"--include-before-body={cover}",
                                      "--embed-resources", "--metadata=pagetitle:Dossier technique"])
    fix_table_layout(page)
    subprocess.run(["weasyprint", str(page), str(OUT / f"{BASENAME}.pdf")], check=True)

    # DOCX: pandoc with a restyled reference document, then tables, cover and a filled table of contents
    ref = BUILD / "reference.docx"
    subprocess.run([pypandoc.get_pandoc_path(), "-o", str(ref), "--print-default-data-file", "reference.docx"],
                   check=True)
    style_reference_docx(ref)
    docx = OUT / f"{BASENAME}.docx"
    pypandoc.convert_file(str(md), "docx", format=PANDOC_FORMAT, outputfile=str(docx),
                          extra_args=["--toc", "--toc-depth=2", "--columns=100000", f"--reference-doc={ref}"])
    style_docx_tables(docx)
    center_docx_images(docx)
    add_docx_cover(docx)
    filled = BUILD / f"{BASENAME}-toc.docx"
    subprocess.run(["/usr/bin/python3", str(TOOLS / "update_toc.py"), str(docx), str(filled)],
                   check=True, capture_output=True, timeout=300)
    shutil.move(str(filled), str(docx))
    print("built", *(p.name for p in sorted(OUT.glob(f"{BASENAME}.*"))))


if __name__ == "__main__":
    main()
